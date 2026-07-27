#!/usr/bin/env python3
"""Backlog Issue Anomaly Scanner — Excel output.

Scans a GitHub repository (default opensourceways/backlog) for issues stuck at
any point in their lifecycle, then writes an Excel report aggregated by submitter
(issue author). Submitter → real-name mapping comes from
opensourceways/opensourceway/community/user-info.yaml.

Sheet 1 "异常统计(按提交人)": one row per submitter, count per anomaly type.
Sheet 2 "异常issue全集": one row per flagged issue (submitter, name, link, types).

Checks dependencies up front: `gh` logged in, and `openpyxl` installed.

Auth via GITHUB_TOKEN / GH_TOKEN / `gh auth token`.
"""

import argparse
import json
import os
import re
import subprocess
import sys
import urllib.error
import urllib.request
from datetime import datetime, timezone

API = "https://api.github.com/graphql"
USER_INFO_URL = (
    "https://raw.githubusercontent.com/opensourceways/opensourceway/"
    "master/community/user-info.yaml"
)

# --- scenario labels (short, stable) ---
S1 = "1-待分类"
S2 = "2-待分配"
S3 = "3-未排期"
S4 = "4-未验收"
S5 = "5-开发停滞"
S6 = "6-PR卡住"
S7 = "7-异常关闭"
S8 = "8-跳过准入"
S9 = "9-rejected未关"
ALL_SCENARIOS = [S1, S2, S3, S4, S5, S6, S7, S8, S9]


def die(msg, code=1):
    print(f"error: {msg}", file=sys.stderr)
    sys.exit(code)


# --------------------------------------------------------------------------
# Dependency checks
# --------------------------------------------------------------------------

def check_gh():
    """Ensure the gh CLI is installed and authenticated. Returns a token."""
    for var in ("GITHUB_TOKEN", "GH_TOKEN"):
        v = os.environ.get(var)
        if v:
            return v, "(from env)"
    if not shutil_which("gh"):
        die("gh CLI not found. Install: https://cli.github.com/  then run `gh auth login`.")
    try:
        r = subprocess.run(["gh", "auth", "token"], capture_output=True, text=True, check=True)
    except subprocess.CalledProcessError:
        die("gh is not logged in. Run: gh auth login")
    tok = r.stdout.strip()
    if not tok:
        die("gh is not logged in. Run: gh auth login")
    return tok, "(from gh auth token)"


def shutil_which(cmd):
    paths = os.environ.get("PATH", "").split(os.pathsep)
    for p in paths:
        candidate = os.path.join(p, cmd)
        if os.path.isfile(candidate) and os.access(candidate, os.X_OK):
            return candidate
    return None


def check_openpyxl():
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        die("openpyxl not installed. Install with:  pip install openpyxl  (or:  pip3 install openpyxl)")
    return openpyxl


# --------------------------------------------------------------------------
# GitHub API
# --------------------------------------------------------------------------

def gql(token, query, variables):
    req = urllib.request.Request(
        API,
        data=json.dumps({"query": query, "variables": variables}).encode("utf-8"),
        headers={
            "Authorization": f"bearer {token}",
            "Content-Type": "application/json",
            "User-Agent": "backlog-issue-anomaly/1.0",
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=60) as resp:
            body = resp.read().decode("utf-8")
    except urllib.error.HTTPError as e:
        die(f"GitHub API HTTP {e.code}: {e.read().decode('utf-8', 'replace')[:500]}")
    except urllib.error.URLError as e:
        die(f"network error: {e}")
    data = json.loads(body)
    if data.get("errors"):
        die(f"GraphQL errors: {json.dumps(data['errors'], ensure_ascii=False)[:800]}")
    return data["data"]


QUERY = """
query($owner:String!, $name:String!, $cursor:String, $states:[IssueState!]) {
  repository(owner:$owner, name:$name) {
    issues(first:100, after:$cursor, states:$states, orderBy:{field:CREATED_AT, direction:DESC}) {
      nodes {
        number title url state createdAt closedAt
        author { login }
        assignees(first:10){nodes{login}}
        milestone{number title}
        labels(first:30){nodes{name}}
        timelineItems(first:100, itemTypes:[LABELED_EVENT, CROSS_REFERENCED_EVENT]) {
          nodes {
            __typename
            ...on LabeledEvent { createdAt label { name } }
            ...on CrossReferencedEvent {
              createdAt
              source {
                __typename
                ...on PullRequest { number state merged mergedAt updatedAt url title }
              }
            }
          }
        }
      }
      pageInfo { endCursor hasNextPage }
    }
  }
}
"""


def fetch_issues(token, owner, name, states):
    issues = []
    cursor = None
    pages = 0
    while True:
        data = gql(token, QUERY, {"owner": owner, "name": name, "cursor": cursor, "states": states})
        repo = data.get("repository")
        if not repo:
            die(f"repository {owner}/{name} not found or no access")
        conn = repo["issues"]
        issues.extend(conn["nodes"])
        pages += 1
        if not conn["pageInfo"]["hasNextPage"]:
            break
        cursor = conn["pageInfo"]["endCursor"]
        if pages > 50:
            break
    return issues


def fetch_user_info(token):
    """Fetch user-info.yaml and return {github_login: name}."""
    req = urllib.request.Request(
        USER_INFO_URL,
        headers={
            "Authorization": f"bearer {token}",
            "User-Agent": "backlog-issue-anomaly/1.0",
        },
    )
    try:
        with urllib.request.urlopen(req, timeout=60) as resp:
            text = resp.read().decode("utf-8")
    except urllib.error.HTTPError as e:
        die(f"failed to fetch user-info.yaml: HTTP {e.code}")
    except urllib.error.URLError as e:
        die(f"failed to fetch user-info.yaml: {e}")
    mapping = {}
    cur_id = None
    for line in text.splitlines():
        m = re.match(r"\s*-\s*github_id:\s*(\S+)", line)
        if m:
            cur_id = m.group(1)
            continue
        m = re.match(r"\s*name:\s*(.+?)\s*$", line)
        if m and cur_id:
            mapping[cur_id] = m.group(1).strip().strip('"').strip("'")
            cur_id = None
    return mapping


# --------------------------------------------------------------------------
# Issue summarization & classification
# --------------------------------------------------------------------------

def parse_dt(s):
    if not s:
        return None
    return datetime.fromisoformat(s.replace("Z", "+00:00"))


def now():
    return datetime.now(timezone.utc)


def days_since(dt, ref=None):
    if dt is None:
        return None
    return ((ref or now()) - dt).total_seconds() / 86400.0


def _nodes(container):
    if not container:
        return []
    return container.get("nodes", []) if isinstance(container, dict) else container


def summarize_issue(it):
    labels = {l["name"] for l in _nodes(it.get("labels"))}
    assignees = [a["login"] for a in _nodes(it.get("assignees"))]
    milestone = it.get("milestone")
    ms = milestone["title"] if milestone else None
    created = parse_dt(it.get("createdAt"))
    author = (it.get("author") or {}).get("login") or "ghost"
    accept_time = None
    prs = []
    for ev in _nodes((it.get("timelineItems") or {}).get("nodes")):
        if ev["__typename"] == "LabeledEvent":
            if (ev.get("label") or {}).get("name") == "accepted" and accept_time is None:
                accept_time = parse_dt(ev.get("createdAt"))
        elif ev["__typename"] == "CrossReferencedEvent":
            src = ev.get("source") or {}
            if src.get("__typename") == "PullRequest":
                prs.append({
                    "number": src.get("number"),
                    "state": src.get("state"),
                    "merged": bool(src.get("merged")),
                    "mergedAt": parse_dt(src.get("mergedAt")),
                    "updatedAt": parse_dt(src.get("updatedAt")),
                    "url": src.get("url"),
                    "title": src.get("title"),
                })
    has_merged_pr = any(p["merged"] for p in prs)
    open_prs = [p for p in prs if p["state"] == "OPEN"]
    return {
        "number": it["number"],
        "title": it.get("title", ""),
        "url": it.get("url", ""),
        "state": it.get("state"),
        "author": author,
        "labels": labels,
        "assignees": assignees,
        "milestone": ms,
        "created": created,
        "accept_time": accept_time,
        "prs": prs,
        "has_merged_pr": has_merged_pr,
        "open_prs": open_prs,
    }


def classify(s, days):
    triggered = []
    is_open = s["state"] == "OPEN"
    accepted = "accepted" in s["labels"]
    rejected = "rejected" in s["labels"]
    resolved_labels = {"wontfix", "duplicate", "invalid", "rejected"}
    closed_as_resolved = bool(s["labels"] & resolved_labels)
    created_age = days_since(s["created"])
    accept_age = days_since(s["accept_time"]) if s["accept_time"] else created_age

    if is_open and not accepted and not rejected and created_age is not None and created_age > days:
        triggered.append(S1)
    if is_open and accepted and not s["assignees"] and accept_age is not None and accept_age > days:
        triggered.append(S2)
    if is_open and accepted and not s["milestone"] and accept_age is not None and accept_age > days:
        triggered.append(S3)
    if is_open and s["has_merged_pr"]:
        triggered.append(S4)
    if is_open and accepted and s["assignees"] and s["milestone"] and not s["prs"] \
            and accept_age is not None and accept_age > days:
        triggered.append(S5)
    stale_open = [p for p in s["open_prs"] if days_since(p["updatedAt"]) is not None
                 and days_since(p["updatedAt"]) > days]
    if is_open and stale_open:
        triggered.append(S6)
    if (not is_open) and not s["has_merged_pr"] and not closed_as_resolved:
        triggered.append(S7)
    if s["has_merged_pr"] and not accepted:
        triggered.append(S8)
    if is_open and rejected:
        triggered.append(S9)
    return triggered


# --------------------------------------------------------------------------
# Excel output
# --------------------------------------------------------------------------

def fmt_age(d):
    if d is None:
        return "?"
    return f"{d:.0f}d" if d >= 1 else f"{int(d * 24)}h"


def write_excel(flagged, user_map, path, days, scanned_open, scanned_closed, repo):
    """flagged: list of (scenario, summary). user_map: login->name."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()

    # ---- Sheet 1: per-submitter counts ----
    ws1 = wb.active
    ws1.title = "异常统计(按提交人)"
    header = ["提交人", "姓名"] + ALL_SCENARIOS + ["合计"]
    ws1.append(header)
    for c in ws1[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
        c.alignment = Alignment(horizontal="center", vertical="center")

    # aggregate
    by_user = {}
    for scn, s in flagged:
        login = s["author"]
        rec = by_user.setdefault(login, {"name": user_map.get(login, login),
                                         "counts": {sc: 0 for sc in ALL_SCENARIOS}})
        rec["counts"][scn] += 1

    rows_sorted = sorted(by_user.items(),
                         key=lambda kv: sum(kv[1]["counts"].values()), reverse=True)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for login, rec in rows_sorted:
        row = [login, rec["name"]] + [rec["counts"][sc] for sc in ALL_SCENARIOS]
        row.append(sum(rec["counts"].values()))
        ws1.append(row)
    # totals row
    totals = [sum(rec["counts"][sc] for _, rec in rows_sorted) for sc in ALL_SCENARIOS]
    ws1.append(["合计", ""] + totals + [sum(totals)])
    last = ws1.max_row
    for c in ws1[last]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")

    # widths & freeze
    ws1.column_dimensions["A"].width = 22
    ws1.column_dimensions["B"].width = 14
    for i in range(len(ALL_SCENARIOS)):
        ws1.column_dimensions[get_column_letter(3 + i)].width = 12
    ws1.column_dimensions[get_column_letter(2 + len(ALL_SCENARIOS) + 1)].width = 10
    ws1.freeze_panes = "C2"

    # ---- Sheet 2: all flagged issues ----
    ws2 = wb.create_sheet("异常issue全集")
    ws2.append(["提交人", "姓名", "Issue链接", "Issue编号", "标题", "异常类型"])
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
        c.alignment = Alignment(horizontal="center", vertical="center")
    # one row per flagged issue, types joined
    seen = {}
    for scn, s in flagged:
        n = s["number"]
        seen.setdefault(n, {"s": s, "types": []})["types"].append(scn)
    for n in sorted(seen.keys(), reverse=True):
        s = seen[n]["s"]
        login = s["author"]
        name = user_map.get(login, login)
        ws2.append([login, name, s["url"], n, s["title"], ", ".join(seen[n]["types"])])
    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 14
    ws2.column_dimensions["C"].width = 52
    ws2.column_dimensions["D"].width = 10
    ws2.column_dimensions["E"].width = 60
    ws2.column_dimensions["F"].width = 28
    ws2.freeze_panes = "A2"
    for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row):
        for c in row:
            c.border = border
            c.alignment = Alignment(vertical="center", wrap_text=(c.column == 5))

    # ---- a small meta sheet at the front ----
    ws0 = wb.create_sheet("说明", 0)
    meta = [
        ["Backlog Issue Anomaly Report"],
        [""],
        ["仓库", repo],
        ["时间窗(天)", days],
        ["扫描 issue 数", scanned_open + scanned_closed],
        ["  其中 open", scanned_open],
        ["  其中 closed", scanned_closed],
        ["标记 issue 数(去重)", len(seen)],
        ["异常类型实例数(含重复)", len(flagged)],
        ["提交人数(有异常)", len(by_user)],
        ["生成时间(UTC)", now().strftime("%Y-%m-%d %H:%M:%S")],
        [""],
        ["异常类型说明"],
        ["1-待分类", "open & 未 accept & 创建超过 N 天"],
        ["2-待分配", "open & accepted & 无 assignee & accept 超过 N 天"],
        ["3-未排期", "open & accepted & 无 milestone & accept 超过 N 天"],
        ["4-未验收", "open & 关联 PR 已合入但 issue 仍 open"],
        ["5-开发停滞", "open & accepted & 有 assignee+milestone & 无 PR & accept 超过 N 天"],
        ["6-PR卡住", "open & 关联 PR 仍 open 且 stale(update 超过 N 天)"],
        ["7-异常关闭", "closed & 无已合入 PR & 非 wontfix/duplicate/invalid/rejected"],
        ["8-跳过准入", "关联 PR 已合入 & issue 从未 accept"],
        ["9-rejected未关", "open & 带 rejected 标签"],
        [""],
        ["注", "backlog 跟踪的交付常在其他仓(om-datacenter/xihe 等);未在 issue 内回链的跨仓 PR 检测不到,故 5/7 为上界。"],
    ]
    for r in meta:
        ws0.append(r)
    ws0["A1"].font = Font(bold=True, size=14)
    ws0.column_dimensions["A"].width = 26
    ws0.column_dimensions["B"].width = 70

    wb.save(path)


# --------------------------------------------------------------------------
# main
# --------------------------------------------------------------------------

def parse_repo(repo):
    if repo.count("/") != 1:
        die(f"--repo must be owner/name, got: {repo}")
    return repo.split("/", 1)


def main():
    ap = argparse.ArgumentParser(description="Scan a GitHub backlog repo for stuck/anomalous issues; output Excel.")
    ap.add_argument("--repo", default="opensourceways/backlog", help="owner/name (default opensourceways/backlog)")
    ap.add_argument("--days", type=int, default=7, help="staleness window in days (default 7)")
    ap.add_argument("--closed", action="store_true", help="also scan closed issues (enables scenarios 7 & 8)")
    ap.add_argument("--output", "-o", default="backlog_anomaly_report.xlsx", help="output .xlsx path")
    args = ap.parse_args()

    print("checking dependencies...", file=sys.stderr)
    check_openpyxl()
    token, src = check_gh()
    print(f"  gh auth: ok {src}", file=sys.stderr)

    owner, name = parse_repo(args.repo)

    print("fetching user-info.yaml ...", file=sys.stderr)
    user_map = fetch_user_info(token)
    print(f"  got {len(user_map)} user mappings", file=sys.stderr)

    states = ["OPEN"] + (["CLOSED"] if args.closed else [])
    print(f"fetching {states} issues from {owner}/{name} ...", file=sys.stderr)
    raw = fetch_issues(token, owner, name, states)
    print(f"  got {len(raw)} issues", file=sys.stderr)

    flagged = []
    scanned_open = 0
    scanned_closed = 0
    for node in raw:
        s = summarize_issue(node)
        if s["state"] == "OPEN":
            scanned_open += 1
        else:
            scanned_closed += 1
        for sc in classify(s, args.days):
            flagged.append((sc, s))

    print(f"writing {args.output} ...", file=sys.stderr)
    write_excel(flagged, user_map, args.output, args.days, scanned_open, scanned_closed, args.repo)
    # summary to stderr
    from collections import Counter
    cnt = Counter(sc for sc, _ in flagged)
    print("done. summary:", file=sys.stderr)
    for sc in ALL_SCENARIOS:
        print(f"  {sc}: {cnt.get(sc, 0)}", file=sys.stderr)
    print(f"output: {args.output}", file=sys.stderr)


if __name__ == "__main__":
    main()
